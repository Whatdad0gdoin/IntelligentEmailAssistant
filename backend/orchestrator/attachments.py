"""Attachment analysis (FR-01/FR-02 extension).

Until this module existed, `_bodies()` in adapters/headers.py threw every
attached part away, so an email whose content was a PDF invoice summarised as
if it were empty. The adapter now hands the bytes over; this file decides what
can be read from them and what the model is allowed to see.

WHAT THIS DOES, AND WHAT IT DELIBERATELY DOES NOT
-------------------------------------------------
Documents (PDF, DOCX, plain text, CSV, Markdown, JSON) have their text
extracted *deterministically*, here, with no model involved. Extracted text is
appended to the prompt and -- this is the part that matters -- to the grounding
source, so a summary that cites a figure from an invoice is verified against
that invoice rather than flagged as an invention.

Images are NOT interpreted. Nothing in this build describes what a picture
shows: the attachment is reported, and the user opens it and looks at it
themselves. An earlier version sent images to a vision model and was removed
deliberately. Two reasons it is not missed. It cost a model call per image,
which against a $5/week cap and a 5 second latency target (NFR-01) is a poor
trade for something the user can see at a glance. And a description of a
picture has no text source behind it, so the grounding layer -- the thing this
project's credibility rests on -- could not check it. Showing the image is
both cheaper and more honest than describing it.

OPTIONAL DEPENDENCIES
---------------------
pypdf and python-docx are imported lazily and their absence is survivable, the
same bargain grounding.py strikes with spaCy. A missing library downgrades one
file type to "cannot read this" and says so in the UI and in EXTRACTOR_BACKENDS
-- it never crashes a summary and never silently pretends the file was empty.

NFR-03
------
Attachment bytes arrive on the SourceEmail, are read in memory here, and are
dropped when the request ends. Nothing is written to disk, nothing is cached,
and every log line in this file carries counts and MIME types only -- never a
filename's contents and never the extracted text.

The viewing route (routes/attachment.py) streams those same in-memory bytes to
the browser on request. That is a display path, not a storage one: the file is
re-read from the mailbox each time and still never touches disk, which is the
same posture the email body has always had.
"""

import csv
import io
import json
import logging

log = logging.getLogger(__name__)


class AttachmentNotFound(RuntimeError):
    """No attachment with that id on that message.

    Not an LLMError: no model is involved anywhere in this file, and inheriting
    from one would put it in the orchestrator-failure bucket in support.py,
    where it would be reported as an AI problem rather than a missing file.
    """

# How much text one attachment may contribute to a prompt. A 200-page PDF
# would otherwise consume the entire token budget and push the email itself
# out of the context window -- the attachment is supporting evidence for the
# message, not a replacement for it.
MAX_ATTACHMENT_TEXT_CHARS = 4000

# How much text all attachments together may contribute.
MAX_TOTAL_ATTACHMENT_CHARS = 12000

# Pages past this are not read. Text-heavy PDFs hit the character cap long
# before this; the page cap is what stops a 900-page scan burning seconds of
# CPU to produce nothing (NFR-01).
MAX_PDF_PAGES = 30

# Statuses, as reported to the UI. Each one is a different thing to tell the
# user, which is why there is no single "failed".
STATUS_EXTRACTED = "extracted"      # text was read
STATUS_EMPTY = "empty"              # read fine, but there was no text in it
STATUS_IMAGE = "image"              # a picture: open it and look at it
STATUS_TOO_LARGE = "too_large"      # bytes were dropped at the adapter
STATUS_UNSUPPORTED = "unsupported"  # nothing in this build can read this type
STATUS_NO_EXTRACTOR = "no_extractor"  # readable in principle, library missing
STATUS_ENCRYPTED = "encrypted"      # password-protected
STATUS_FAILED = "failed"            # corrupt or malformed

# Types the browser may render INLINE, in an iframe or an <img>, when the user
# opens an attachment.
#
# This list is a security boundary, not a convenience. An attachment is bytes
# chosen by whoever sent the email, served from the app's own origin, so
# anything the browser will execute as markup in that origin is a stored XSS
# vector against the signed-in user. text/html is the obvious one. image/svg+xml
# is the non-obvious one -- an SVG is a document that can carry <script>, and
# treating it as "just an image" is a classic way to ship this bug.
#
# Everything absent from this list is still downloadable. It is served as a
# plain byte stream with an attachment disposition, so the browser saves it
# instead of running it.
INLINE_SAFE_TYPES = frozenset({
    "image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp",
    "application/pdf",
    "text/plain", "text/csv",
})

_PLAIN_TEXT_TYPES = frozenset({
    "text/plain", "text/markdown", "text/x-markdown", "text/csv",
    "text/tab-separated-values", "application/json", "text/rtf",
})

_PDF_TYPE = "application/pdf"
_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Which extractor actually ran, for the same reason grounding.py records its
# entity backend: an evaluation run must be able to say what produced its
# numbers. Filled in on first use of each type.
EXTRACTOR_BACKENDS = {}


class ExtractedAttachment:
    """What one attachment turned out to contain.

    `text` is extracted content, capped. `status` says why it is empty when it
    is -- the difference between "this PDF had no text layer" and "pypdf is not
    installed" matters to the person reading the UI, and collapsing both into
    an empty string is how an install problem gets mistaken for an empty file.
    """

    __slots__ = ("id", "filename", "content_type", "size_bytes", "status",
                 "text", "truncated", "detail")

    def __init__(self, attachment, status, text="", truncated=False, detail=""):
        self.id = attachment.id
        self.filename = attachment.filename
        self.content_type = attachment.content_type
        self.size_bytes = attachment.size_bytes
        self.status = status
        self.text = text
        self.truncated = truncated
        self.detail = detail

    @property
    def has_text(self):
        return bool(self.text.strip())

    @property
    def can_preview(self):
        """True when the browser may safely render this inline.

        False does not mean the file is unavailable -- it means opening it
        downloads it rather than displaying it. See INLINE_SAFE_TYPES.
        """
        return (
            self.can_open
            and self.content_type.lower().split(";")[0].strip() in INLINE_SAFE_TYPES
        )

    @property
    def can_open(self):
        """True when there are bytes to hand back at all."""
        return self.status != STATUS_TOO_LARGE and self.size_bytes > 0

    def as_api_dict(self):
        """The shape the frontend renders. Never includes the bytes."""
        return {
            "id": self.id,
            "filename": self.filename,
            "content_type": self.content_type,
            "size_bytes": self.size_bytes,
            "status": self.status,
            "has_text": self.has_text,
            "truncated": self.truncated,
            # Whether the user can open it, and whether opening it shows the
            # file or saves it. Two separate questions: a .zip can be opened
            # but not previewed.
            "can_open": self.can_open,
            "can_preview": self.can_preview,
            "detail": self.detail,
        }


# --- extractors ------------------------------------------------------------


def _truncate(text, limit=MAX_ATTACHMENT_TEXT_CHARS):
    text = (text or "").strip()
    if len(text) <= limit:
        return text, False
    window = text[:limit]
    boundary = window.rfind("\n")
    if boundary < limit * 0.6:
        boundary = window.rfind(" ")
    if boundary > limit * 0.6:
        window = window[:boundary]
    return window.rstrip(), True


def _decode_text(payload):
    """Decode bytes that are supposed to be text, without guessing wildly."""
    for encoding in ("utf-8", "utf-16", "cp1252"):
        try:
            return payload.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return payload.decode("utf-8", errors="replace")


def _extract_plain(payload, content_type):
    text = _decode_text(payload)
    if content_type == "application/json":
        # Pretty-print so the model sees structure rather than one long line.
        # A JSON attachment that will not parse is still text worth reading, so
        # a parse failure falls through to the raw decode rather than erroring.
        try:
            text = json.dumps(json.loads(text), indent=2)[:MAX_ATTACHMENT_TEXT_CHARS * 2]
        except (ValueError, TypeError):
            pass
    elif content_type in ("text/csv", "text/tab-separated-values"):
        text = _format_table(text, delimiter="\t" if "tab" in content_type else ",")
    return text


def _format_table(raw, delimiter=","):
    """Render a CSV as aligned rows.

    A spreadsheet pasted in as raw comma-separated text reads badly to a model
    and worse to the grounding checker, which would see "1200,Due,2026-03-01"
    as one token soup. Rows are capped: the point is to let the model say what
    the file is about, not to feed it a database.
    """
    try:
        rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    except (csv.Error, ValueError):
        return raw
    if not rows:
        return ""
    kept = rows[:60]
    lines = [" | ".join(str(cell).strip() for cell in row) for row in kept if any(row)]
    if len(rows) > len(kept):
        lines.append(f"... and {len(rows) - len(kept)} more rows")
    return "\n".join(lines)


def _extract_pdf(payload):
    """Text layer of a PDF. Returns (status, text, detail)."""
    try:
        import pypdf
    except ImportError:
        EXTRACTOR_BACKENDS["pdf"] = "missing"
        log.info(
            "pypdf is not installed; PDF attachments cannot be read. "
            "Install with: pip install -r backend/requirements.txt"
        )
        return STATUS_NO_EXTRACTOR, "", "PDF reading is not installed on the server."

    EXTRACTOR_BACKENDS["pdf"] = "pypdf"
    try:
        reader = pypdf.PdfReader(io.BytesIO(payload))
        if reader.is_encrypted:
            # An empty-password PDF is common and legitimately readable.
            try:
                if not reader.decrypt(""):
                    return STATUS_ENCRYPTED, "", "This PDF is password protected."
            except Exception:
                return STATUS_ENCRYPTED, "", "This PDF is password protected."

        pages = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:
                # One bad page should not lose the other twenty-nine.
                log.warning("pdf: skipping a page (%s)", type(exc).__name__)
            if sum(len(p) for p in pages) > MAX_ATTACHMENT_TEXT_CHARS * 2:
                break
    except Exception as exc:
        log.warning("pdf: could not be read (%s)", type(exc).__name__)
        return STATUS_FAILED, "", "This PDF could not be read."

    text = "\n".join(pages).strip()
    if not text:
        # Overwhelmingly a scan: an image of a page, with no text layer. Saying
        # so is more useful than "empty", because the two have different fixes.
        return (STATUS_EMPTY, "",
                "This PDF has no text layer -- it is most likely a scan.")
    return STATUS_EXTRACTED, text, ""


def _extract_docx(payload):
    """Paragraph and table text from a .docx. Returns (status, text, detail)."""
    try:
        import docx
    except ImportError:
        EXTRACTOR_BACKENDS["docx"] = "missing"
        log.info(
            "python-docx is not installed; Word attachments cannot be read. "
            "Install with: pip install -r backend/requirements.txt"
        )
        return STATUS_NO_EXTRACTOR, "", "Word reading is not installed on the server."

    EXTRACTOR_BACKENDS["docx"] = "python-docx"
    try:
        document = docx.Document(io.BytesIO(payload))
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    except Exception as exc:
        log.warning("docx: could not be read (%s)", type(exc).__name__)
        return STATUS_FAILED, "", "This Word document could not be read."

    text = "\n".join(blocks).strip()
    if not text:
        return STATUS_EMPTY, "", "This document contains no text."
    return STATUS_EXTRACTED, text, ""


# --- entry points ----------------------------------------------------------


def extract_one(attachment):
    """Read one Attachment. Never raises: a bad file is a status, not a crash."""
    content_type = (attachment.content_type or "").lower().split(";")[0].strip()

    if attachment.too_large:
        return ExtractedAttachment(
            attachment, STATUS_TOO_LARGE,
            detail="This file is too large to analyse.",
        )

    if not attachment.data:
        return ExtractedAttachment(attachment, STATUS_EMPTY, detail="This file is empty.")

    if content_type.startswith("image/"):
        # No text to extract and nothing here interprets pictures. The status
        # exists so the UI can say "an image, open it to see it" rather than
        # reporting a readable file that happened to yield nothing.
        return ExtractedAttachment(attachment, STATUS_IMAGE)

    try:
        if content_type == _PDF_TYPE:
            status, text, detail = _extract_pdf(attachment.data)
        elif content_type == _DOCX_TYPE:
            status, text, detail = _extract_docx(attachment.data)
        elif content_type in _PLAIN_TEXT_TYPES or content_type.startswith("text/"):
            text = _extract_plain(attachment.data, content_type)
            status = STATUS_EXTRACTED if text.strip() else STATUS_EMPTY
            detail = "" if text.strip() else "This file is empty."
        else:
            return ExtractedAttachment(
                attachment, STATUS_UNSUPPORTED,
                detail=f"{content_type or 'This file type'} cannot be read.",
            )
    except Exception as exc:
        # Defensive: an extractor that raises something unforeseen must not
        # take down the summary of the email it was attached to.
        log.warning("attachment: extractor crashed on %s (%s)",
                    content_type, type(exc).__name__)
        return ExtractedAttachment(attachment, STATUS_FAILED,
                                   detail="This file could not be read.")

    text, truncated = _truncate(text)
    log.info(
        "attachment %s: %s, %d bytes in, %d chars out, truncated=%s",
        attachment.id[:8], content_type, attachment.size_bytes, len(text), truncated,
    )
    return ExtractedAttachment(attachment, status, text=text, truncated=truncated,
                               detail=detail)


def find(email, attachment_id):
    """The named attachment on this message, or None.

    Ids are content-addressed and computed per message (headers.py), so an id
    from one email cannot select an attachment on another -- which is what
    stops the viewing route becoming a way to read files off other people's
    mail by guessing identifiers.

    Searches the visible attachments only: an inline layout image is never
    offered in the UI, so a request naming one did not come from the UI.
    """
    for attachment in visible(getattr(email, "attachments", None)):
        if attachment.id == attachment_id:
            return attachment
    return None


def visible(attachments):
    """The attachments a person would call attachments.

    Drops the inline images that HTML mail uses as layout -- logos, spacers,
    tracking pixels (see Attachment.is_embedded). Applied everywhere the
    count or the list is surfaced, so the inbox paperclip, the reading pane
    and the summary response all agree on what is attached.
    """
    return [a for a in (attachments or []) if not a.is_embedded]


def extract_all(attachments):
    """Read every attachment on a message, within the shared character budget.

    The budget is spent in order. An email with one readable PDF and nine
    unreadable images spends nothing on the images, so the PDF still gets its
    full allowance.
    """
    extracted = []
    spent = 0
    for attachment in visible(attachments):
        item = extract_one(attachment)
        if item.has_text:
            remaining = MAX_TOTAL_ATTACHMENT_CHARS - spent
            if remaining <= 0:
                item.text = ""
                item.truncated = True
            elif len(item.text) > remaining:
                item.text = item.text[:remaining].rstrip()
                item.truncated = True
            spent += len(item.text)
        extracted.append(item)
    return extracted


def prompt_block(extracted):
    """The attachment section of a prompt, or "" when there is nothing to say.

    Both the readable and the unreadable attachments are listed. Telling the
    model that an unreadable file is present is the difference between a
    summary that says "a scanned invoice is attached" and one that does not
    mention the invoice at all -- and the second is the one that misleads.
    """
    if not extracted:
        return ""

    lines = []
    for item in extracted:
        size_kb = max(1, round(item.size_bytes / 1024))
        header = f'<attachment filename="{item.filename}" type="{item.content_type}" size="{size_kb} KB">'
        if item.has_text:
            body = item.text
            if item.truncated:
                body += "\n[... truncated]"
        elif item.status == STATUS_IMAGE:
            # The model is told the picture exists and told not to guess at
            # it. Nothing in this build looks at images, and a model invited
            # to speculate from a filename will describe one anyway.
            body = ("[An image. Its contents are not available to you. "
                    "Mention that an image is attached; do not describe it.]")
        else:
            body = f"[{item.detail or 'Content could not be read.'}]"
        lines.append(f"{header}\n{body}\n</attachment>")

    return (
        "\n\nThe email has the following attachment(s). Their content is part "
        "of the message and may be summarised, but say when something is "
        "attached and unreadable rather than ignoring it.\n\n"
        + "\n\n".join(lines)
    )


def filename_list(attachments):
    """Attachment names and sizes on one line, for the classifier prompt.

    Deliberately not the extracted text. Classification runs as one batch over
    a whole inbox, so adding a page of PDF text per email would multiply the
    prompt many times over to improve a four-way label that a filename usually
    settles on its own.
    """
    parts = []
    for attachment in visible(attachments):
        size_kb = max(1, round(attachment.size_bytes / 1024))
        parts.append(f"{attachment.filename} ({attachment.content_type}, {size_kb} KB)")
    return ", ".join(parts)


def grounding_text(extracted):
    """Attachment text as grounding source material.

    Without this, a summary citing a total from an attached invoice is flagged
    as an invented number: the figure is real, it is simply not in the email
    body. Filenames are included because a summary naming the attached file is
    making a supportable claim about the message.
    """
    parts = []
    for item in extracted or []:
        parts.append(item.filename)
        if item.has_text:
            parts.append(item.text)
    return "\n".join(parts)
